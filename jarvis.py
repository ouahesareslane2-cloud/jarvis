#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
JARVIS v2 : assistant personnel intégré à Windows, façon Tony Stark (se fige en un seul jarvis.exe).
  Il comprend toutes les langues (darija comprise), répond à la voix, voit l’écran, pilote souris et clavier,
  administre Windows (processus, services, registre, tâches, réseau, périphériques), gère fichiers, archives,
  images, PowerPoint, recherches web, scripts, Git, bases de données, et réalise des projets entiers seul.
  Sans clé IA : toutes les commandes rapides. Avec une clé gratuite Gemini ou Groq (bouton « Clé IA ») : l’agent complet.
Raccourci Ctrl+Alt+J pour le rappeler. Dis « démarre avec Windows » pour l’auto-lancement.
"""
import os, sys, re, ast, json, time, math, queue, random, shutil, socket, locale, operator, datetime, threading, subprocess, unicodedata, webbrowser, collections
import base64, hashlib, tempfile, zipfile, html as htm
import urllib.request, urllib.parse, urllib.error

# ───────────── RÉGLAGES ─────────────
NOM = "jarvis"                 # mot d'éveil
LANGUE = "fr"                  # fr, en ou ar
VILLE = ""                     # vide = détectée automatiquement pour la météo

APPS = {"chrome": "chrome", "google chrome": "chrome", "firefox": "firefox", "edge": "msedge",
        "bloc-notes": "notepad", "bloc notes": "notepad", "notepad": "notepad", "calculatrice": "calc",
        "explorateur": "explorer", "fichiers": "explorer", "paint": "mspaint", "parametres": "ms-settings:",
        "gestionnaire des taches": "taskmgr", "terminal": "cmd", "word": "winword", "excel": "excel",
        "powerpoint": "powerpnt", "spotify": "spotify", "discord": "discord", "vscode": "code", "steam": "steam"}
SITES = {"youtube": "https://www.youtube.com", "google": "https://www.google.com", "gmail": "https://mail.google.com",
         "maps": "https://maps.google.com", "wikipedia": "https://fr.wikipedia.org", "facebook": "https://facebook.com",
         "instagram": "https://instagram.com", "netflix": "https://www.netflix.com", "github": "https://github.com"}

STT = {"fr": "fr-FR", "en": "en-US", "ar": "ar-DZ"}[LANGUE]
LANGUES_VOIX = ["fr-FR", "ar-DZ", "en-US"]  # langues reconnues automatiquement : français, darija, anglais…
CULTURE = {"fr": "fr-FR", "en": "en-US", "ar": "ar-SA"}[LANGUE]
DOSSIER = os.path.join(os.path.expanduser("~"), ".jarvis")
FICHIER = os.path.join(DOSSIER, "notes.json")
os.makedirs(DOSSIER, exist_ok=True)
try: sys.stdout.reconfigure(encoding="utf-8")
except Exception: pass
try: locale.setlocale(locale.LC_TIME, "")
except Exception: pass
os.system("")  # active les couleurs ANSI sous Windows

try:
    import numpy as np, sounddevice as sd, speech_recognition as sr
    MICRO = True
except Exception:
    MICRO = False
try: import psutil
except Exception: psutil = None

try:
    import tkinter as tk
    from tkinter import simpledialog, messagebox
except Exception:
    tk = None
ui, verrou, CONFIG, modele_cache = queue.Queue(), threading.Lock(), {}, {}
historique, confirmation, SEUIL = [], None, 500
FLAGS = {"ecoute": MICRO, "direct": False}
FROZEN, HOME = getattr(sys, "frozen", False), os.path.expanduser("~")

_Popen = subprocess.Popen
class _PopenSilencieux(_Popen):  # pas de fenêtre noire, et pas d’erreur « handle invalide » dans un .exe sans console
    def __init__(self, *a, **k):
        if os.name == "nt":
            k["creationflags"] = k.get("creationflags", 0) | 0x08000000
            for flux in ("stdin", "stdout", "stderr"): k.setdefault(flux, subprocess.DEVNULL)
        super().__init__(*a, **k)
subprocess.Popen = _PopenSilencieux

# ───────────── OUTILS ─────────────
def http_json(url, data=None, timeout=10, headers=None):
    req = urllib.request.Request(url, json.dumps(data).encode() if data is not None else None, {"Content-Type": "application/json", "User-Agent": "Jarvis", **(headers or {})})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return json.loads(r.read().decode("utf-8"))

def sans_accents(s):
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")

def notes():
    try:
        with open(FICHIER, encoding="utf-8") as f: return json.load(f)
    except Exception: return []

def sauver(liste):
    with open(FICHIER, "w", encoding="utf-8") as f: json.dump(liste, f, ensure_ascii=False, indent=1)

def dire_sapi(texte):
    texte = re.sub(r"[*_#`>]", "", texte)
    try:
        if os.name == "nt":
            ps = ("Add-Type -AssemblyName System.Speech;$s=New-Object System.Speech.Synthesis.SpeechSynthesizer;"
                  "try{$s.SelectVoiceByHints('NotSet','NotSet',0,[cultureinfo]'" + CULTURE + "')}catch{};$s.Speak($env:JARVIS_TXT)")
            subprocess.run(["powershell", "-NoProfile", "-Command", ps], env={**os.environ, "JARVIS_TXT": texte}, creationflags=0x08000000)
        else:
            subprocess.run(["say" if sys.platform == "darwin" else "espeak", texte])
    except Exception:
        pass

def parler(texte):
    print(f"\033[96mJarvis ▸\033[0m {texte}")
    ui.put(("msg", "ai", texte)); ui.put(("etat", "speak"))
    dire(texte)
    ui.put(("etat", "idle"))

def bip():
    try:
        import winsound
        for _ in range(3): winsound.Beep(880, 300)
    except Exception:
        print("\a")

def touche(vk, fois=1):  # touches multimédia Windows : 0xAD muet, 0xAE vol-, 0xAF vol+, 0xB0 suivant, 0xB1 précédent, 0xB3 lecture
    if os.name != "nt": return False
    import ctypes
    for _ in range(fois):
        ctypes.windll.user32.keybd_event(vk, 0, 0, 0)
        ctypes.windll.user32.keybd_event(vk, 0, 2, 0)
    return True

def media(vk, fois, ok):
    return ok if touche(vk, fois) else "Cette action ne marche que sous Windows."

def win(commande_os, ok):
    if os.name != "nt": return "Cette action ne marche que sous Windows."
    subprocess.Popen(commande_os, shell=True)
    return ok

def lancer(cible):
    cible = cible.replace('"', "")
    if os.name == "nt": subprocess.Popen(f'start "" "{cible}"', shell=True)
    else: subprocess.Popen(["open" if sys.platform == "darwin" else "xdg-open", cible])

def planifier(sec, msg):
    t = threading.Timer(sec, lambda: (bip(), parler("Rappel : " + msg)))
    t.daemon = True
    t.start()

OPS = {ast.Add: operator.add, ast.Sub: operator.sub, ast.Mult: operator.mul, ast.Div: operator.truediv,
       ast.Pow: operator.pow, ast.Mod: operator.mod, ast.USub: operator.neg}

def evaluer(expr):
    def ev(n):
        if isinstance(n, ast.Constant) and isinstance(n.value, (int, float)): return n.value
        if isinstance(n, ast.BinOp) and type(n.op) in OPS:
            a, b = ev(n.left), ev(n.right)
            if isinstance(n.op, ast.Pow) and abs(b) > 100: raise ValueError
            return OPS[type(n.op)](a, b)
        if isinstance(n, ast.UnaryOp) and type(n.op) in OPS: return OPS[type(n.op)](ev(n.operand))
        raise ValueError
    return ev(ast.parse(expr.strip(), mode="eval").body)

def capture():
    try:
        from PIL import ImageGrab
        chemin = os.path.join(os.path.expanduser("~"), "Pictures", f"jarvis_{datetime.datetime.now():%Y%m%d_%H%M%S}.png")
        os.makedirs(os.path.dirname(chemin), exist_ok=True)
        ImageGrab.grab().save(chemin)
        return "Capture enregistrée dans ton dossier Images."
    except ImportError:
        return "Installe le module pillow pour les captures d’écran."

def infos():
    if not psutil: return "Installe le module psutil pour que je lise l’état du PC."
    s = f"Processeur à {psutil.cpu_percent(interval=0.5):.0f} pour cent, mémoire à {psutil.virtual_memory().percent:.0f} pour cent."
    b = psutil.sensors_battery()
    if b: s += f" Batterie à {b.percent:.0f} pour cent{', en charge' if b.power_plugged else ''}."
    return s

def meteo(ville):
    try:
        url = "https://wttr.in/" + urllib.parse.quote(ville) + "?format=%25l:+%25C,+%25t&lang=" + LANGUE
        with urllib.request.urlopen(urllib.request.Request(url, headers={"User-Agent": "curl"}), timeout=8) as r:
            return r.read().decode("utf-8").strip()
    except Exception:
        return "Je n’arrive pas à joindre le service météo."

DOSSIERS = {"telechargements": "Downloads", "downloads": "Downloads", "documents": "Documents", "bureau": "Desktop", "desktop": "Desktop",
            "images": "Pictures", "photos": "Pictures", "musique": "Music", "videos": "Videos"}
BLAGUES = ["Pourquoi les plongeurs plongent-ils toujours en arrière ? Parce que sinon ils tombent dans le bateau.",
           "Qu’est-ce qu’un crocodile qui surveille la pharmacie ? Un Lacoste garde.",
           "Je voulais te raconter une blague sur le Wi-Fi, mais elle ne passe pas bien."]
AIDE = ("Je sais ouvrir applis, sites et dossiers, fermer une appli, régler volume et luminosité, piloter la musique, faire des captures, "
        "verrouiller, mettre en veille ou éteindre le PC, gérer notes, minuteurs et rappels à heure fixe, calculer, donner la météo, "
        "retrouver un fichier, dicter du texte dans n’importe quelle fenêtre, lire l’état du PC. Avec une clé IA gratuite, je comprends toutes les langues, même la darija, et je réalise des projets entiers : recherches, fichiers, images, PowerPoint, scripts, administration Windows.")

def quitter():
    parler("À bientôt.")
    os._exit(0)

def combo(msg, *vks):  # appuie sur un raccourci clavier Windows (ex. Win+D)
    if os.name != "nt": return "Cette action ne marche que sous Windows."
    import ctypes
    k = ctypes.windll.user32.keybd_event
    for v in vks: k(v, 0, 0, 0)
    for v in reversed(vks): k(v, 0, 2, 0)
    return msg

def trouver(nom):
    cible, fin = sans_accents(nom).lower().replace(" ", ""), time.time() + 20
    for dossier in ("Desktop", "Documents", "Downloads", "Pictures", "Music", "Videos"):
        for racine, _, fichiers in os.walk(os.path.join(os.path.expanduser("~"), dossier)):
            if time.time() > fin: return "La recherche prend trop de temps. Donne-moi un nom plus précis."
            for f in fichiers:
                if cible in sans_accents(f).lower().replace(" ", ""):
                    lancer(os.path.join(racine, f)); return f"J’ouvre {f}."
    return f"Je n’ai trouvé aucun fichier « {nom} » dans tes dossiers personnels."

def wiki(sujet):
    try:
        base = f"https://{LANGUE}.wikipedia.org"
        titre = http_json(base + "/w/api.php?action=query&list=search&format=json&srlimit=1&srsearch=" + urllib.parse.quote(sujet))["query"]["search"][0]["title"]
        texte = http_json(base + "/api/rest_v1/page/summary/" + urllib.parse.quote(titre.replace(" ", "_"), safe=""))["extract"]
        return " ".join(re.split(r"(?<=[.!?]) ", texte)[:3])
    except Exception:
        return None

def gourmands():
    if not psutil: return "Installe le module psutil pour que je lise les processus."
    p = [x for x in psutil.process_iter(["name", "memory_info"]) if x.info["memory_info"]]
    p = sorted(p, key=lambda x: x.info["memory_info"].rss, reverse=True)[:3]
    return "Les plus gourmands en mémoire : " + ", ".join(f"{x.info['name']} ({x.info['memory_info'].rss / 1e9:.1f} giga)" for x in p) + "."

def briefing():
    h, nb = datetime.datetime.now(), len(notes())
    s = f"Bonjour ! Il est {h:%H:%M}, nous sommes le {h:%A %d %B}. " + meteo(VILLE).rstrip(".") + ". "
    if psutil and psutil.sensors_battery(): s += f"Batterie à {psutil.sensors_battery().percent:.0f} pour cent. "
    return s + (f"Tu as {nb} note{'s' if nb > 1 else ''} en attente." if nb else "Aucune note en attente.")

def demarrage_auto():
    if os.name != "nt": return "Cette action ne marche que sous Windows."
    dossier = os.path.join(os.environ["APPDATA"], "Microsoft", "Windows", "Start Menu", "Programs", "Startup")
    pyw = os.path.join(os.path.dirname(sys.executable), "pythonw.exe")
    lancement = f'"{sys.executable}"' if FROZEN else f'"{pyw}" "{os.path.abspath(__file__)}"'
    with open(os.path.join(dossier, "Jarvis.bat"), "w", encoding="utf-8") as f:
        f.write(f'@echo off\nchcp 65001 >nul\nstart "" {lancement}\n')
    return "C’est fait : je me lancerai tout seul au démarrage de Windows."

# ───────────── COMMANDES LOCALES (instantanées) ─────────────
def commande(texte):
    """Renvoie la réponse si c'est une commande PC, sinon None (la question va au cerveau)."""
    global confirmation
    o = unicodedata.normalize("NFC", texte.strip())
    n = sans_accents(o).lower().replace("’", "'")
    maintenant = datetime.datetime.now()

    if confirmation:
        action, confirmation = confirmation, None
        return action() if re.search(r"\b(oui|ok|confirme|yes|vas-y)\b", n) else "D’accord, j’annule."
    if re.match(r"(au revoir|a plus|a bientot|bye|goodbye|quitte|ferme-toi)", n):
        quitter()
    if re.search(r"quelle heure|l'heure|what time", n): return f"Il est {maintenant:%H:%M}."
    if re.search(r"quel jour|la date|what day", n): return f"Nous sommes le {maintenant:%A %d %B %Y}."

    m = re.search(r"(?:calcule|calculer|combien (?:font|fait|vaut)|calculate)\s+(.+)", n)
    if m:
        e = re.sub(r"\bfois\b|\bx\b|×", "*", m.group(1))
        e = re.sub(r"\bplus\b", "+", e)
        e = re.sub(r"\bmoins\b", "-", e)
        e = re.sub(r"divise par|÷|\bsur\b", "/", e).replace(",", ".")
        e = re.sub(r"[^0-9+\-*/().% ]", "", e)
        try:
            r = round(evaluer(e), 6)
            return f"{e.strip()} égale {int(r) if r == int(r) else r}."
        except Exception:
            return "Je n’arrive pas à calculer ça."

    m = re.match(r"(?:note|notez|prends note|retiens|rappelle-toi)\b\s*(?:que|:|de)?\s*(.+)", n)
    if m:
        liste = notes(); liste.append(o[m.start(1):]); sauver(liste)
        return f"Noté. Tu as {len(liste)} note{'s' if len(liste) > 1 else ''}."
    if re.search(r"mes notes|lis.*notes|liste.*notes", n):
        liste = notes()
        return ("Tes notes : " + " ; ".join(f"{i + 1}, {x}" for i, x in enumerate(liste))) if liste else "Tu n’as aucune note."
    if re.search(r"(efface|supprime|vide).*notes", n):
        sauver([]); return "Toutes tes notes sont effacées."

    m = re.search(r"(?:volume|\bson)\s*(?:a|to|at)?\s*(\d{1,3})\b", n)
    if m:
        v = min(100, int(m.group(1)))
        touche(0xAE, 50); touche(0xAF, round(v / 2))
        return f"Volume à {v} pour cent."
    m = re.search(r"luminosite\D*(\d{1,3})", n)
    if m:
        v = min(100, int(m.group(1)))
        return win(f'powershell -NoProfile -Command "(Get-WmiObject -Namespace root/WMI -Class WmiMonitorBrightnessMethods).WmiSetBrightness(1,{v})"', f"Luminosité à {v} pour cent. Ça ne marche que sur un écran d’ordinateur portable.")
    m = re.search(r"(?:rappelle|reveille)\S*\s.*?\ba\s*(\d{1,2})\s*(?:h|:)\s*(\d{2})?", n)
    if m:
        cible = maintenant.replace(hour=int(m.group(1)) % 24, minute=int(m.group(2) or 0), second=0)
        if cible <= maintenant: cible += datetime.timedelta(days=1)
        msg = re.sub(r"^(?:de |pour |que |d')+", "", o[m.end():].strip()) or "C’est l’heure."
        planifier((cible - maintenant).total_seconds(), msg)
        return f"Rappel programmé à {cible:%H:%M}."
    m = re.match(r"(?:ferme|fermer|tue|kill|close)\s+(?:le |la |l'|les )?(.+)", n)
    if m:
        c = m.group(1).strip()
        if re.fullmatch(r"fenetre|window", c): return combo("Fenêtre fermée.", 0x12, 0x73)
        exe = re.sub(r"[^\w.-]", "", APPS.get(c, c).replace(" ", ""))
        return win(f"taskkill /f /im {exe}.exe", f"Je ferme {c}.")
    if re.search(r"affiche (?:le )?bureau|minimise tout|reduis tout|show desktop", n): return combo("Bureau affiché.", 0x5B, 0x44)
    if re.search(r"change de fenetre|fenetre suivante|switch window", n): return combo("Fenêtre suivante.", 0x12, 0x09)
    if re.search(r"plein ecran|full ?screen", n): return combo("C’est fait.", 0x7A)
    m = re.match(r"(?:dicte|ecris|tape|type)\s+(.+)", n)
    if m:
        if os.name != "nt": return "La dictée ne marche que sous Windows."
        ui.put(("cacher",)); time.sleep(0.6)
        subprocess.run(["powershell", "-NoProfile", "-Command", "Set-Clipboard -Value $env:JARVIS_TXT"], env={**os.environ, "JARVIS_TXT": o[m.start(1):]}, creationflags=0x08000000)
        return combo("J’écris.", 0x11, 0x56)
    m = re.match(r"(?:ouvre|ouvrir|affiche)\s+(?:mes |le |la |les |mon )?(?:dossier )?(telechargements|downloads|documents|bureau|desktop|images|photos|musique|videos)\b", n)
    if m:
        lancer(os.path.join(os.path.expanduser("~"), DOSSIERS[m.group(1)])); return f"J’ouvre {m.group(1)}."
    m = re.match(r"(?:trouve|cherche|retrouve)\s+(?:le |la |mon |ma |un |une )?(?:fichier|document|pdf)\s+(.+)", n)
    if m: return trouver(m.group(1).strip())
    m = re.match(r"(?:c'est quoi|qu'est-ce que|qu'est ce que|qui est|qui etait|parle-moi de|definis|what is|who is)\s+(?:un |une |le |la |les |l')?(.+)", n)
    if m and not CONFIG.get("cle"):
        rep = wiki(o[m.start(1):])
        if rep: return rep
    if re.search(r"blague|fais-moi rire|joke", n): return random.choice(BLAGUES)
    if re.search(r"adresse ip|\bmon ip\b", n):
        sk = socket.socket(socket.AF_INET, socket.SOCK_DGRAM); sk.connect(("8.8.8.8", 80))
        return "Ton adresse IP locale est " + sk.getsockname()[0] + "."
    if re.search(r"espace disque|stockage|disque dur", n):
        u = shutil.disk_usage(os.path.expanduser("~"))
        return f"Il reste {u.free / 1e9:.0f} gigaoctets libres sur {u.total / 1e9:.0f}."
    if re.search(r"(?:vide|nettoie).*corbeille", n): return win('powershell -NoProfile -Command "Clear-RecycleBin -Force"', "Corbeille vidée.")
    if re.search(r"ralentit|gourmand|processus|plus de memoire", n): return gourmands()
    if re.search(r"demarre.*(?:avec windows|au demarrage)|lance-toi au demarrage|demarrage automatique", n): return demarrage_auto()
    if re.match(r"(?:bonjour|salut|coucou)\b", n) or re.search(r"briefing|resume de ma journee|quoi de neuf", n): return briefing()
    if re.match(r"(?:aide|help|que sais-tu faire|tes fonctions|tes commandes)\b", n): return AIDE

    d = re.search(r"(\d+)\s*(secondes?|sec|minutes?|min|heures?|h)\b", n)
    if d and re.search(r"minuteur|chrono|timer|rappelle|reveille|alarme", n):
        v, u = int(d.group(1)), d.group(2)[0]
        msg = re.sub(r"^(?:de |pour |que |d')+", "", o[d.end():].strip()) or "Le minuteur est terminé."
        planifier(v * (3600 if u == "h" else 60 if u == "m" else 1), msg)
        return f"C’est parti : {v} {d.group(2)}."

    if re.search(r"volume|\bson\b|sound", n):
        if re.search(r"monte|augmente|plus fort|\bup\b|raise", n): return media(0xAF, 8, "Volume augmenté.")
        if re.search(r"baisse|diminue|moins fort|down|lower", n): return media(0xAE, 8, "Volume baissé.")
        if re.search(r"coupe|muet|mute|silence", n): return media(0xAD, 1, "Son coupé.")
    if re.search(r"\bpause\b|^play\b|reprends la (?:musique|lecture)", n): return media(0xB3, 1, "C’est fait.")
    if re.search(r"(?:suivant|suivante|next)", n) and re.search(r"musique|piste|chanson|morceau|next", n): return media(0xB0, 1, "Piste suivante.")
    if re.search(r"precedent|previous", n): return media(0xB1, 1, "Piste précédente.")

    if re.search(r"capture|screenshot", n): return capture()
    if re.search(r"verrouill|\block\b", n): return win("rundll32.exe user32.dll,LockWorkStation", "Session verrouillée.")
    if re.search(r"\bveille\b|\bsleep\b", n): return win("rundll32.exe powrprof.dll,SetSuspendState 0,1,0", "Mise en veille.")
    if re.search(r"annule.*(extinction|arret|redemarrage)", n): return win("shutdown /a", "Extinction annulée.")
    if re.search(r"eteins|eteindre|shutdown|arrete (?:l'ordinateur|le pc)", n):
        confirmation = lambda: win("shutdown /s /t 20", "Extinction dans vingt secondes. Dis « annule l’extinction » pour arrêter.")
        return "Tu confirmes l’extinction de l’ordinateur ?"
    if re.search(r"redemarre|reboot|restart", n):
        confirmation = lambda: win("shutdown /r /t 20", "Redémarrage dans vingt secondes. Dis « annule l’extinction » pour arrêter.")
        return "Tu confirmes le redémarrage de l’ordinateur ?"

    if re.search(r"batterie|processeur|\bcpu\b|\bram\b|memoire|etat du (?:systeme|pc)|performances?", n): return infos()
    if re.search(r"meteo|quel temps|weather", n):
        v = re.search(r"\b(?:a|de|pour|in|sur)\s+([a-z' -]+?)\s*$", n)
        return meteo(v.group(1) if v else VILLE)

    m = re.match(r"(?:ouvre|ouvrir|lance|lancer|demarre|open|launch|start)\s+(?:le |la |l'|les |un |une |mon |ma )?(.+)", n)
    if m:
        c = m.group(1).strip()
        if c in SITES: webbrowser.open(SITES[c]); return f"J’ouvre {c}."
        if c in APPS: lancer(APPS[c]); return f"J’ouvre {c}."
        if re.search(r"\.\w{2,}$", c): webbrowser.open("https://" + c.replace(" ", "")); return f"J’ouvre {c}."
        lancer(c); return f"J’essaie d’ouvrir {c}."
    m = re.match(r"(?:cherche|recherche|search|google|trouve(?:-moi)?)\s+(?:sur (?:google|internet|le web) )?(.+)", n)
    if m:
        q = o[m.start(1):]
        webbrowser.open("https://www.google.com/search?q=" + urllib.parse.quote_plus(q)); return f"Je cherche {q}."
    m = re.match(r"(?:joue|mets|mets-moi|youtube)\s+(.+)", n)
    if m:
        q = o[m.start(1):]
        webbrowser.open("https://www.youtube.com/results?search_query=" + urllib.parse.quote_plus(q)); return f"Je lance YouTube : {q}."
    return None

# ───────────── CERVEAU (optionnel : clé IA gratuite Gemini ou Groq) ─────────────
FOURNISSEURS = (("AIza", "https://generativelanguage.googleapis.com/v1beta/openai"), ("gsk_", "https://api.groq.com/openai/v1"))
OUTIL = [{"type": "function", "function": {"name": "commande_pc", "parameters": {"type": "object", "properties": {"ordre": {"type": "string"}}, "required": ["ordre"]},
          "description": "Exécute une action sur le PC de l’utilisateur à partir d’un ordre court en français. Exemples : « ouvre chrome », « ouvre mes téléchargements », "
                         "« ferme spotify », « volume à 30 », « luminosité 60 », « minuteur 5 minutes », « rappelle-moi à 18h de sortir », « note que… », « mes notes », "
                         "« cherche … sur google », « joue … sur youtube », « capture d’écran », « météo à Oran », « trouve le fichier … », « écris … » (tape le texte dans la fenêtre active), "
                         "« verrouille le PC », « éteins le PC », « état du système », « affiche le bureau »."}}]

def charger_config():
    try:
        with open(os.path.join(DOSSIER, "config.json"), encoding="utf-8") as f: CONFIG.update(json.load(f))
    except Exception:
        pass
    if not CONFIG.get("cle"): CONFIG["cle"] = os.environ.get("JARVIS_CLE", "")

def sauver_cle(cle):
    CONFIG["cle"] = cle.strip(); modele_cache.clear()
    with open(os.path.join(DOSSIER, "config.json"), "w", encoding="utf-8") as f: json.dump(CONFIG, f)

def choisir_modele(base, cle):
    if "m" not in modele_cache:
        gem = "generativelanguage" in base
        try:
            ids = [x["id"].replace("models/", "") for x in http_json(base + "/models", timeout=15, headers={"Authorization": "Bearer " + cle})["data"]]
            if gem: c = sorted((i for i in ids if re.fullmatch(r"gemini-[\d.]+-flash", i)), key=lambda i: float(re.search(r"[\d.]+", i).group()), reverse=True)
            else: c = [i for p in (r"llama-3\.3-70b", r"70b", r"llama") for i in ids if re.search(p, i)]
            modele_cache["m"] = (c or ids)[0]
        except Exception:
            modele_cache["m"] = "gemini-2.5-flash" if gem else "llama-3.3-70b-versatile"
    return modele_cache["m"]

# ───────────── AGENT : OUTILS (ce que Jarvis peut faire sur le PC) ─────────────
REG, IMAGES, UI_OK = {}, [], [False]
DANGER = re.compile(r"\b(?:format-volume|diskpart|clear-disk|remove-partition|bcdedit|cipher\s+/w)\b|\bformat\s+[a-z]:|(?:remove-item|\brd\b|\brmdir\b|\bdel\b|\berase\b)[^\n]*?(?:\b[a-z]:\\?(?:\s|$)|\\windows|\\system32|\\users\\?(?:\s|$))", re.I)
LECTURE = re.compile(r"\s*(?:get-\S+|select(?:-object)?|where(?:-object)?|sort(?:-object)?|measure(?:-object)?|group(?:-object)?|format-(?:table|list)|ft|fl|out-string|test-(?:connection|path|netconnection)|resolve-dnsname|ipconfig|ping|tracert|nslookup|netstat|tasklist|systeminfo|whoami|hostname|dir|ls|type|cat|findstr|find|echo|nvidia-smi|git\s+(?:status|log|diff|branch|remote|show)|reg\s+query|sc\s+query|schtasks\s+/query|netsh\s+(?:wlan|interface)\s+show|wmic\s+\S+\s+get)\b", re.I)
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Jarvis/2.0"}
CATEGORIES = {"Images": ".jpg .jpeg .png .gif .webp .bmp .svg .heic", "Documents": ".pdf .doc .docx .txt .odt .rtf .xls .xlsx .csv .ppt .pptx",
              "Archives": ".zip .rar .7z .tar .gz", "Installateurs": ".exe .msi .msix", "Videos": ".mp4 .mkv .avi .mov .webm",
              "Musique": ".mp3 .wav .flac .m4a .ogg", "Code": ".py .js .html .css .json .xml .bat .ps1"}
JOURNAL_TRI = os.path.join(DOSSIER, "dernier_tri.json")
VK = {"ctrl": 0x11, "alt": 0x12, "shift": 0x10, "win": 0x5B, "enter": 0x0D, "tab": 0x09, "esc": 0x1B, "space": 0x20, "backspace": 0x08,
      "suppr": 0x2E, "delete": 0x2E, "haut": 0x26, "bas": 0x28, "gauche": 0x25, "droite": 0x27, "debut": 0x24, "fin": 0x23}

def outil(desc, **params):
    def deco(f):
        REG[f.__name__] = (f, desc, params)
        return f
    return deco

def schema_outils():
    sortie = []
    for nom, (f, desc, p) in REG.items():
        fn = {"name": nom, "description": desc, "parameters": {"type": "object", "properties": {k: {"type": "string", "description": v.lstrip("?")} for k, v in p.items()}}}
        req = [k for k, v in p.items() if not v.startswith("?")]
        if req: fn["parameters"]["required"] = req
        sortie.append({"type": "function", "function": fn})
    return sortie

def executer_outil(nom, args):
    if nom not in REG: return "Outil inconnu : " + nom
    f, _, p = REG[nom]
    try:
        return str(f(**{k: v for k, v in args.items() if k in p}))[:8000]
    except TypeError as e:
        return f"Arguments invalides : {e}"
    except Exception as e:
        return f"Erreur : {type(e).__name__} : {e}"

def est_lecture_seule(cmd):
    if re.search(r"[<>]|-outfile|\biex\b|invoke-expression|invoke-webrequest|invoke-restmethod|start-process|\bcurl\b|\bwget\b|\$\(|`", cmd, re.I): return False
    parties = [x for x in re.split(r"[|;&\n]+", cmd) if x.strip()]
    return bool(parties) and all(LECTURE.match(x) for x in parties)

def confirmer(texte):
    if tk is None or "--texte" in sys.argv:
        return input(texte + " [o/N] ").strip().lower().startswith(("o", "y"))
    res, evt = {}, threading.Event()
    ui.put(("confirmer", texte, res, evt))
    evt.wait(600)
    return bool(res.get("ok"))

def autoriser(detail):
    danger = bool(DANGER.search(detail))
    if CONFIG.get("confiance") and not danger: return True
    return confirmer("Jarvis veut exécuter :\n\n" + detail[:900] + ("\n\n⚠ ACTION POTENTIELLEMENT DESTRUCTRICE" if danger else "") + "\n\nAutoriser ?")

def projet_courant(): return CONFIG.get("projet") or os.path.join(HOME, "Jarvis", "Projets", "General")
def chemin_abs(c): return os.path.join(projet_courant(), os.path.expanduser(c))
def nettoyer(h): return htm.unescape(re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", h))).strip()
def vk(nom):
    nom = nom.strip().lower()
    if nom in VK: return VK[nom]
    m = re.fullmatch(r"f(\d{1,2})", nom)
    return 0x6F + int(m.group(1)) if m else ord(nom.upper()[0])

def shell(args, timeout=120, env=None):
    r = subprocess.run(args, capture_output=True, timeout=timeout, env=env)
    t = ((r.stdout or b"").decode("utf-8", "replace") + (r.stderr or b"").decode("utf-8", "replace")).strip()
    return t[:6000] or f"(terminé, code {r.returncode})"

def http_brut(url, donnees=None, methode="GET", entetes=None, timeout=60):
    req = urllib.request.Request(url, donnees, {**UA, **(entetes or {})}, method=methode)
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return r.status, r.read()

def corbeille(p):
    if os.name != "nt":
        shutil.rmtree(p) if os.path.isdir(p) else os.remove(p)
        return
    m = "DeleteDirectory" if os.path.isdir(p) else "DeleteFile"
    shell(["powershell", "-NoProfile", "-Command", f"Add-Type -AssemblyName Microsoft.VisualBasic; [Microsoft.VisualBasic.FileIO.FileSystem]::{m}($env:JARVIS_P,'OnlyErrorDialogs','SendToRecycleBin')"], env={**os.environ, "JARVIS_P": p})

def contient(p, t):
    with open(p, encoding="utf-8", errors="ignore") as f: return t.lower() in f.read().lower()

# ── commandes rapides, PowerShell, CMD, code ──
@outil("Exécute un ordre simple du routeur de commandes rapides du PC (ouvrir applis, sites et dossiers, volume, médias, notes, minuteurs, rappels, météo…).", ordre="ordre court en français, ex : « ouvre chrome », « volume à 30 », « minuteur 5 minutes », « note que… », « météo à Oran »")
def commande_pc(ordre): return commande(ordre) or "Commande non reconnue : utilise un autre outil."

@outil("Exécute du PowerShell sur le PC : processus, services, registre, tâches planifiées, réseau, périphériques, variables d’environnement, alimentation, git, winget, FFmpeg… Renvoie la sortie.", script="code PowerShell")
def powershell(script):
    if os.name != "nt": return "PowerShell n’est disponible que sous Windows."
    if not est_lecture_seule(script) and not autoriser("PowerShell :\n" + script): return "Refusé par l’utilisateur."
    return shell(["powershell", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-Command", "[Console]::OutputEncoding=[Text.Encoding]::UTF8; " + script])

@outil("Exécute une commande CMD (invite de commandes) ou un script .bat.", ligne="ligne de commande CMD")
def cmd(ligne):
    if os.name != "nt": return "CMD n’est disponible que sous Windows."
    if not est_lecture_seule(ligne) and not autoriser("CMD :\n" + ligne): return "Refusé par l’utilisateur."
    return shell(["cmd", "/c", "chcp 65001>nul & " + ligne])

@outil("Écrit puis exécute du code et renvoie sa sortie. langage : python (interpréteur intégré : Pillow, python-pptx, pypdf, openpyxl, psutil, sqlite3, csv, json, zipfile…), python_systeme (Python installé sur le PC, pour pip), javascript (Node), powershell, batch.", langage="python, python_systeme, javascript, powershell ou batch", code="le code complet")
def executer_code(langage, code):
    l = langage.lower().strip()
    ext = {"python": ".py", "python_systeme": ".py", "javascript": ".js", "powershell": ".ps1", "batch": ".bat"}.get(l)
    if not ext: return "Langage inconnu."
    if not autoriser(f"Code {l} :\n{code}"): return "Refusé par l’utilisateur."
    dossier = projet_courant(); os.makedirs(dossier, exist_ok=True)
    tmp = tempfile.mkdtemp(prefix="jarvis_"); f = os.path.join(tmp, "script" + ext); sortie = os.path.join(tmp, "sortie.txt")
    with open(f, "w", encoding="utf-8-sig" if ext == ".ps1" else "utf-8") as fh: fh.write(("@echo off\nchcp 65001>nul\n" + code) if ext == ".bat" else code)
    if l == "python":
        subprocess.run(([sys.executable] if FROZEN else [sys.executable, os.path.abspath(__file__)]) + ["--exec", f, sortie], timeout=300, cwd=dossier)
        with open(sortie, encoding="utf-8", errors="replace") as fh: return fh.read()[:6000] or "(aucune sortie)"
    cmds = {"python_systeme": [shutil.which("py") or shutil.which("python") or "python", f], "javascript": ["node", f],
            "powershell": ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", f], "batch": ["cmd", "/c", f]}
    r = subprocess.run(cmds[l], capture_output=True, timeout=300, cwd=dossier)
    return ((r.stdout or b"").decode("utf-8", "replace") + (r.stderr or b"").decode("utf-8", "replace")).strip()[:6000] or "(aucune sortie)"

@outil("Relance Jarvis avec les droits administrateur (nécessaire pour le registre HKLM, les services, les périphériques, certaines tâches système).")
def relancer_admin():
    if os.name != "nt": return "Windows seulement."
    import ctypes
    if ctypes.windll.shell32.IsUserAnAdmin(): return "Je suis déjà administrateur."
    if not confirmer("Relancer Jarvis en administrateur ? Windows te demandera ton accord."): return "Refusé par l’utilisateur."
    if ctypes.windll.shell32.ShellExecuteW(None, "runas", sys.executable, "" if FROZEN else f'"{os.path.abspath(__file__)}"', None, 1) > 32: os._exit(0)
    return "Élévation refusée ou impossible."

# ── fichiers ──
@outil("Crée (si besoin) et sélectionne le dossier d’un projet ; les chemins relatifs des autres outils pointent ensuite dedans.", nom="nom du projet")
def projet_dossier(nom):
    CONFIG["projet"] = os.path.join(HOME, "Jarvis", "Projets", re.sub(r"[^\w.-]+", "_", nom).strip("_") or "Projet")
    os.makedirs(CONFIG["projet"], exist_ok=True)
    return "Dossier du projet : " + CONFIG["projet"]

@outil("Lit un fichier (texte, code, json, csv, docx, pdf, xlsx) et renvoie son contenu, tronqué à 8000 caractères.", chemin="chemin du fichier", debut="?position de départ en caractères")
def lire_fichier(chemin, debut="0"):
    p, e = chemin_abs(chemin), os.path.splitext(chemin)[1].lower()
    if e == ".docx":
        with zipfile.ZipFile(p) as z: t = nettoyer(z.read("word/document.xml").decode("utf-8").replace("</w:p>", "\n"))
    elif e == ".pdf":
        from pypdf import PdfReader
        t = "\n".join(pg.extract_text() or "" for pg in PdfReader(p).pages)
    elif e in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        t = "\n".join("\t".join("" if c is None else str(c) for c in r) for ws in load_workbook(p, read_only=True, data_only=True) for r in ws.iter_rows(values_only=True))
    else:
        with open(p, encoding="utf-8", errors="replace") as f: t = f.read()
    d = int(debut or 0)
    return t[d:d + 8000] + (f"\n[… tronqué : {len(t)} caractères au total]" if len(t) > d + 8000 else "")

@outil("Crée, écrase ou complète (mode ajout) un fichier texte ou de code. Un chemin relatif va dans le dossier du projet courant.", chemin="chemin du fichier", contenu="contenu complet", mode="?ecraser (défaut) ou ajout")
def ecrire_fichier(chemin, contenu, mode="ecraser"):
    p = chemin_abs(chemin)
    if os.path.exists(p) and not mode.startswith("aj") and not autoriser("Écraser le fichier :\n" + p): return "Refusé par l’utilisateur."
    os.makedirs(os.path.dirname(p), exist_ok=True)
    with open(p, "a" if mode.startswith("aj") else "w", encoding="utf-8") as f: f.write(contenu)
    return "Fichier écrit : " + p

@outil("Recherche des fichiers sur les disques selon le nom, l’extension, la taille minimale, la date de modification ou le contenu texte. 40 résultats max.", nom="?partie du nom", extension="?ex : pdf", contenu="?texte à trouver dans le fichier", taille_min_mo="?taille minimale en Mo", modifie_apres="?date AAAA-MM-JJ", dossier="?où chercher (défaut : tous les disques)")
def chercher_fichiers(nom="", extension="", contenu="", taille_min_mo="", modifie_apres="", dossier=""):
    racines = [os.path.expanduser(dossier)] if dossier else ([f"{c}:\\" for c in "CDEFGHIJ" if os.path.exists(f"{c}:\\")] if os.name == "nt" else [HOME])
    nom, ext, fin, trouves = sans_accents(nom).lower(), extension.lower().lstrip("."), time.time() + 60, []
    depuis = datetime.datetime.strptime(modifie_apres, "%Y-%m-%d").timestamp() if modifie_apres else 0
    mini = float(taille_min_mo or 0) * 1e6
    for racine in racines:
        for r, dossiers, noms in os.walk(racine):
            dossiers[:] = [d for d in dossiers if d.lower() not in ("windows", "$recycle.bin", "system volume information", "node_modules", ".git", "appdata")]
            if time.time() > fin or len(trouves) >= 40: break
            for f in noms:
                if (nom and nom not in sans_accents(f).lower()) or (ext and not f.lower().endswith("." + ext)): continue
                p = os.path.join(r, f)
                try:
                    st = os.stat(p)
                    if st.st_size < mini or st.st_mtime < depuis: continue
                    if contenu and (st.st_size > 5e6 or not contient(p, contenu)): continue
                except OSError:
                    continue
                trouves.append(f"{p}  ({st.st_size / 1e6:.1f} Mo, {datetime.datetime.fromtimestamp(st.st_mtime):%Y-%m-%d})")
    return "\n".join(trouves) or "Aucun fichier trouvé."

@outil("Gère fichiers et dossiers : creer_dossier, copier, deplacer, supprimer (envoyé à la corbeille).", action="creer_dossier, copier, deplacer ou supprimer", source="chemin source", destination="?chemin de destination")
def fichiers(action, source, destination=""):
    s, d = chemin_abs(source), (chemin_abs(destination) if destination else "")
    if action != "creer_dossier" and not autoriser(f"{action} : {s}" + (f" → {d}" if d else "")): return "Refusé par l’utilisateur."
    if action == "creer_dossier": os.makedirs(s, exist_ok=True)
    elif action == "copier": shutil.copytree(s, d) if os.path.isdir(s) else shutil.copy2(s, d)
    elif action == "deplacer": shutil.move(s, d)
    elif action == "supprimer": corbeille(s)
    else: return "Action inconnue."
    return "Fait."

@outil("Range un dossier en désordre en sous-dossiers par type (Images, Documents, Archives, Installateurs, Videos, Musique, Code, Autres). simulation=oui pour voir le plan sans rien déplacer.", dossier="dossier à ranger, chemin réel (ex : Downloads, Desktop, Documents)", simulation="?oui pour ne rien déplacer")
def trier_dossier(dossier, simulation="non"):
    base, plan = os.path.join(HOME, os.path.expanduser(dossier)), []
    for f in sorted(os.listdir(base)):
        p = os.path.join(base, f)
        if os.path.isfile(p):
            e = os.path.splitext(f)[1].lower()
            plan.append((p, os.path.join(base, next((k for k, v in CATEGORIES.items() if e in v.split()), "Autres"), f)))
    if simulation.lower().startswith(("o", "y")): return "\n".join(f"{os.path.basename(a)} → {os.path.basename(os.path.dirname(b))}" for a, b in plan[:80]) or "Rien à ranger."
    if not plan: return "Rien à ranger."
    if not autoriser(f"Ranger {len(plan)} fichiers de {base} dans des sous-dossiers par type ?"): return "Refusé par l’utilisateur."
    fait = []
    for a, b in plan:
        os.makedirs(os.path.dirname(b), exist_ok=True)
        n, (rac, ext) = 1, os.path.splitext(b)
        while os.path.exists(b): b, n = f"{rac} ({n}){ext}", n + 1
        shutil.move(a, b); fait.append([a, b])
    with open(JOURNAL_TRI, "w", encoding="utf-8") as j: json.dump(fait, j)
    return f"{len(fait)} fichiers rangés. Je peux annuler ce tri si tu veux."

@outil("Annule le dernier rangement fait par trier_dossier.")
def annuler_tri():
    with open(JOURNAL_TRI, encoding="utf-8") as j: fait = json.load(j)
    for a, b in reversed(fait):
        if os.path.exists(b): shutil.move(b, a)
    return f"{len(fait)} fichiers remis à leur place."

@outil("Compresse (zip, tar.gz) ou extrait (zip, tar, tar.gz ; rar et 7z si 7-Zip est installé).", action="compresser ou extraire", source="fichier ou dossier source", destination="?archive à créer ou dossier de destination")
def archive(action, source, destination=""):
    s = chemin_abs(source).rstrip("\\/")
    if action.lower().startswith("comp"):
        d = chemin_abs(destination) if destination else s + ".zip"
        fmt, base = ("gztar" if d.endswith((".tar.gz", ".tgz")) else "zip"), re.sub(r"\.(zip|tar\.gz|tgz)$", "", d)
        return "Archive créée : " + shutil.make_archive(base, fmt, root_dir=os.path.dirname(s), base_dir=os.path.basename(s))
    d = chemin_abs(destination) if destination else os.path.splitext(s)[0]
    try: shutil.unpack_archive(s, d)
    except shutil.ReadError: return shell([shutil.which("7z") or r"C:\Program Files\7-Zip\7z.exe", "x", s, "-o" + d, "-y"])
    return "Extrait dans " + d

@outil("Calcule l’empreinte d’un fichier (md5, sha1, sha256…) et la compare à une valeur attendue.", chemin="fichier", algo="?md5, sha1 ou sha256 (défaut)", attendu="?empreinte attendue à comparer")
def hash_fichier(chemin, algo="sha256", attendu=""):
    h = hashlib.new(algo.lower() or "sha256")
    with open(chemin_abs(chemin), "rb") as f:
        for bloc in iter(lambda: f.read(1 << 20), b""): h.update(bloc)
    d = h.hexdigest()
    return d + (("  ✔ identique" if d.lower() == attendu.strip().lower() else "  ✘ DIFFÉRENT") if attendu else "")

@outil("Nettoie le système : fichiers temporaires de l’utilisateur et corbeille. Renvoie l’espace libéré.")
def nettoyer_systeme():
    if not autoriser("Supprimer les fichiers temporaires (%TEMP%) et vider la corbeille ?"): return "Refusé par l’utilisateur."
    avant, tmp = shutil.disk_usage(HOME).free, tempfile.gettempdir()
    for e in os.listdir(tmp):
        try:
            p = os.path.join(tmp, e)
            shutil.rmtree(p) if os.path.isdir(p) else os.remove(p)
        except Exception:
            pass
    if os.name == "nt": shell(["powershell", "-NoProfile", "-Command", "Clear-RecycleBin -Force -ErrorAction SilentlyContinue"])
    return f"Libéré : {(shutil.disk_usage(HOME).free - avant) / 1e6:.0f} Mo."

# ── système ──
@outil("Surveille le PC en temps réel : CPU, RAM, chaque disque, batterie, GPU NVIDIA et température si disponibles.")
def etat_systeme():
    if not psutil: return "Module psutil manquant."
    vm = psutil.virtual_memory()
    l = [f"CPU : {psutil.cpu_percent(interval=1)} % ({psutil.cpu_count()} cœurs logiques)", f"RAM : {vm.percent} % ({vm.used / 1e9:.1f} sur {vm.total / 1e9:.1f} Go)"]
    for pt in psutil.disk_partitions():
        try:
            u = psutil.disk_usage(pt.mountpoint)
            l.append(f"Disque {pt.device} : {u.free / 1e9:.0f} Go libres sur {u.total / 1e9:.0f} Go")
        except Exception:
            pass
    b = psutil.sensors_battery()
    if b: l.append(f"Batterie : {b.percent:.0f} %{' (en charge)' if b.power_plugged else ''}")
    if shutil.which("nvidia-smi"): l.append("GPU : " + shell(["nvidia-smi", "--query-gpu=name,utilization.gpu,temperature.gpu,memory.used,memory.total", "--format=csv,noheader"]))
    if os.name == "nt":
        t = shell(["powershell", "-NoProfile", "-Command", "Get-CimInstance -Namespace root/wmi -ClassName MSAcpi_ThermalZoneTemperature | % { [math]::Round($_.CurrentTemperature/10-273.15,1) }"])
        if re.fullmatch(r"[\d.,\s]+", t): l.append("Température : " + t.replace("\n", " / ") + " °C")
    return "\n".join(l)

@outil("Gère les processus : lister (les plus gourmands), tuer (de force, par nom ou PID), priorite (basse, normale, haute).", action="lister, tuer ou priorite", cible="?nom du processus ou PID", valeur="?basse, normale ou haute")
def processus(action, cible="", valeur=""):
    if not psutil: return "Module psutil manquant."
    if action.startswith("list"):
        p = sorted(psutil.process_iter(["name", "pid", "memory_info"]), key=lambda x: x.info["memory_info"].rss if x.info["memory_info"] else 0, reverse=True)[:15]
        return "\n".join(f"{x.info['pid']}  {x.info['name']}  {x.info['memory_info'].rss / 1e6:.0f} Mo" for x in p if x.info["memory_info"])
    cibles = [x for x in psutil.process_iter(["name", "pid"]) if str(x.info["pid"]) == cible or (cible and cible.lower().replace(".exe", "") == (x.info["name"] or "").lower().replace(".exe", ""))]
    if not cibles: return "Aucun processus trouvé."
    if not autoriser(f"{action} : {len(cibles)} processus « {cible} »"): return "Refusé par l’utilisateur."
    for x in cibles:
        if action.startswith("tu"): x.kill()
        else: x.nice({"basse": psutil.IDLE_PRIORITY_CLASS, "normale": psutil.NORMAL_PRIORITY_CLASS, "haute": psutil.HIGH_PRIORITY_CLASS}[valeur.lower()])
    return f"{len(cibles)} processus traités."

@outil("Liste les ports réseau en écoute avec le programme qui les utilise, ou libère un port en tuant le processus qui l’occupe.", action="lister ou liberer", port="?numéro de port")
def ports(action="lister", port=""):
    if not psutil: return "Module psutil manquant."
    lignes = []
    for c in psutil.net_connections(kind="inet"):
        if c.status == psutil.CONN_LISTEN and c.laddr and c.pid:
            try: lignes.append((c.laddr.port, c.pid, psutil.Process(c.pid).name()))
            except Exception: pass
    if action.startswith("lib"):
        cibles = {pid for p, pid, _ in lignes if str(p) == str(port)}
        if not cibles: return "Aucun programme n’écoute sur ce port."
        if not autoriser(f"Tuer les processus {sorted(cibles)} qui occupent le port {port}"): return "Refusé par l’utilisateur."
        for pid in cibles: psutil.Process(pid).kill()
        return "Port libéré."
    return "\n".join(f"port {p}  PID {pid}  {n}" for p, pid, n in sorted(set(lignes))[:60]) or "Aucun port en écoute."

@outil("Lit le presse-papiers (texte) ou y place du texte.", action="lire ou ecrire", texte="?texte à placer dans le presse-papiers")
def presse_papiers(action, texte=""):
    if os.name != "nt": return "Windows seulement."
    if action.startswith("ec"):
        shell(["powershell", "-NoProfile", "-Command", "Set-Clipboard -Value $env:JARVIS_TXT"], env={**os.environ, "JARVIS_TXT": texte})
        return "Copié."
    return shell(["powershell", "-NoProfile", "-Command", "[Console]::OutputEncoding=[Text.Encoding]::UTF8; Get-Clipboard -Raw"])

@outil("Ouvre un fichier, un dossier, une application ou un lien web avec le programme par défaut.", cible="chemin, nom d’appli ou URL")
def ouvrir(cible):
    if re.match(r"https?://", cible): webbrowser.open(cible)
    else: lancer(APPS.get(cible.lower(), os.path.expanduser(cible)))
    return "Ouvert."

# ── yeux, souris, clavier ──
@outil("Prend une capture de l’écran, l’enregistre et te la montre (pour voir l’interface, lire du texte à l’écran ou décider où cliquer).")
def capture_ecran():
    import io
    from PIL import ImageGrab
    img = ImageGrab.grab(all_screens=True)
    p = os.path.join(HOME, "Pictures", "Jarvis", f"ecran_{datetime.datetime.now():%Y%m%d_%H%M%S}.png")
    os.makedirs(os.path.dirname(p), exist_ok=True); img.save(p)
    if not CONFIG.get("cle", "").startswith("AIza"): return f"Capture enregistrée : {p} (je ne peux la regarder qu’avec une clé Gemini)."
    L, H = img.size
    img.thumbnail((1600, 1600)); b = io.BytesIO(); img.convert("RGB").save(b, "JPEG", quality=80)
    IMAGES.append(base64.b64encode(b.getvalue()).decode())
    return f"Capture enregistrée : {p}. Écran réel : {L}x{H} px ; image jointe réduite à {img.width}x{img.height} : multiplie les coordonnées de l’image par {L / img.width:.3f} avant de cliquer."

@outil("Contrôle la souris et le clavier : clic, double_clic, clic_droit, deplacer (x, y en pixels de l’écran), defiler (valeur dans texte : positif = haut, négatif = bas), taper (texte à saisir), touches (combinaison, ex : ctrl+c, alt+tab, win+d, enter).", action="clic, double_clic, clic_droit, deplacer, defiler, taper ou touches", x="?abscisse en pixels", y="?ordonnée en pixels", texte="?texte à taper, combinaison de touches ou valeur de défilement")
def souris_clavier(action, x="0", y="0", texte=""):
    if os.name != "nt": return "Windows seulement."
    if not UI_OK[0]:
        if not confirmer("Jarvis veut contrôler la souris et le clavier pour cette demande. Autoriser ?"): return "Refusé par l’utilisateur."
        UI_OK[0] = True
    import ctypes
    u, a = ctypes.windll.user32, action.lower()
    if a in ("clic", "double_clic", "clic_droit", "deplacer"):
        u.SetCursorPos(int(float(x)), int(float(y)))
        if a != "deplacer":
            bas, haut = (0x08, 0x10) if a == "clic_droit" else (0x02, 0x04)
            for _ in range(2 if a == "double_clic" else 1):
                u.mouse_event(bas, 0, 0, 0, 0); u.mouse_event(haut, 0, 0, 0, 0); time.sleep(0.06)
    elif a == "defiler": u.mouse_event(0x0800, 0, 0, (int(float(texte or 3)) * 120) & 0xFFFFFFFF, 0)
    elif a == "taper":
        shell(["powershell", "-NoProfile", "-Command", "Set-Clipboard -Value $env:JARVIS_TXT"], env={**os.environ, "JARVIS_TXT": texte})
        combo("", 0x11, 0x56)
    elif a == "touches": combo("", *[vk(t) for t in texte.split("+")])
    else: return "Action inconnue."
    time.sleep(0.2)
    return "Fait."

# ── web ──
@outil("Recherche sur le web (vraie recherche Google avec une clé Gemini, sinon DuckDuckGo) et renvoie une synthèse ou des résultats avec leurs liens.", requete="la recherche")
def web_recherche(requete):
    cle = CONFIG.get("cle", "")
    if cle.startswith("AIza"):
        try:
            modele = choisir_modele(FOURNISSEURS[0][1], cle)
            r = http_json(f"https://generativelanguage.googleapis.com/v1beta/models/{modele}:generateContent", {"contents": [{"parts": [{"text": requete}]}], "tools": [{"google_search": {}}]}, 90, {"x-goog-api-key": cle})
            c = r["candidates"][0]
            src = [f"{g['web'].get('title', '')} : {g['web']['uri']}" for g in c.get("groundingMetadata", {}).get("groundingChunks", []) if "web" in g]
            return "".join(x.get("text", "") for x in c["content"]["parts"]) + "\nSources :\n" + "\n".join(src[:6])
        except Exception:
            pass
    page = http_brut("https://html.duckduckgo.com/html/?q=" + urllib.parse.quote_plus(requete))[1].decode("utf-8", "replace")
    sortie = []
    for u, t, d in re.findall(r'class="result__a" href="([^"]+)"[^>]*>(.*?)</a>.*?class="result__snippet"[^>]*>(.*?)</a>', page, re.S)[:8]:
        if "uddg=" in u: u = urllib.parse.unquote(u.split("uddg=")[1].split("&")[0])
        sortie.append(f"{nettoyer(t)} : {u}\n   {nettoyer(d)}")
    return "\n".join(sortie) or "Aucun résultat."

@outil("Télécharge une page web et renvoie son texte lisible (tronqué à 7000 caractères). Sert à lire les résultats d’une recherche et à extraire des données.", url="adresse de la page")
def web_lire(url):
    h = http_brut(url)[1].decode("utf-8", "replace")
    t = nettoyer(re.sub(r"(?is)<(script|style|noscript|svg)[^>]*>.*?</\1>", " ", h))
    return t[:7000] + (" […]" if len(t) > 7000 else "")

@outil("Télécharge un ou plusieurs fichiers depuis des liens directs vers un dossier.", urls="liens directs, un par ligne", destination="?dossier de destination (défaut : Telechargements du projet)")
def telecharger(urls, destination=""):
    d, sortie = chemin_abs(destination or "Telechargements"), []
    os.makedirs(d, exist_ok=True)
    for u in re.split(r"[\s,]+", urls.strip()):
        if u:
            p = os.path.join(d, os.path.basename(urllib.parse.urlparse(u).path) or "fichier")
            with open(p, "wb") as f: f.write(http_brut(u, timeout=600)[1])
            sortie.append(p)
    return "Téléchargé :\n" + "\n".join(sortie)

@outil("Envoie une requête HTTP (GET, POST, PUT, DELETE…) vers une API locale ou distante et renvoie le statut et la réponse.", methode="GET, POST…", url="adresse complète", corps="?corps de la requête (JSON ou texte)", entetes="?en-têtes au format JSON")
def requete_http(methode, url, corps="", entetes=""):
    h = json.loads(entetes) if entetes else {}
    if corps.lstrip()[:1] in ("{", "["): h.setdefault("Content-Type", "application/json")
    try: s, b = http_brut(url, corps.encode("utf-8") if corps else None, methode.upper(), h)
    except urllib.error.HTTPError as e: s, b = e.code, e.read()
    return f"HTTP {s}\n" + b.decode("utf-8", "replace")[:4000]

# ── création : images, PowerPoint, audio ──
@outil("Génère une image à partir d’une description détaillée (de préférence en anglais) et l’enregistre ; renvoie le chemin. Sert aussi à illustrer les PowerPoint.", description="description détaillée de l’image", fichier="?chemin du fichier à créer (.png ou .jpg)")
def creer_image(description, fichier=""):
    p = chemin_abs(fichier) if fichier else os.path.join(HOME, "Pictures", "Jarvis", re.sub(r"\W+", "_", description)[:40].strip("_") + ".png")
    os.makedirs(os.path.dirname(p), exist_ok=True)
    cle, donnees = CONFIG.get("cle", ""), None
    if cle.startswith("AIza"):
        try:
            noms = [m["name"].replace("models/", "") for m in http_json("https://generativelanguage.googleapis.com/v1beta/models?pageSize=200", timeout=20, headers={"x-goog-api-key": cle})["models"]]
            mod = sorted((n for n in noms if re.fullmatch(r"gemini-[\d.]+-flash-image", n)), reverse=True) or ["gemini-2.5-flash-image"]
            r = http_json(f"https://generativelanguage.googleapis.com/v1beta/models/{mod[0]}:generateContent", {"contents": [{"parts": [{"text": description}]}], "generationConfig": {"responseModalities": ["TEXT", "IMAGE"]}}, 120, {"x-goog-api-key": cle})
            for part in r["candidates"][0]["content"]["parts"]:
                if "inlineData" in part: donnees = base64.b64decode(part["inlineData"]["data"])
        except Exception:
            donnees = None
    if not donnees:
        donnees = http_brut("https://image.pollinations.ai/prompt/" + urllib.parse.quote(description) + f"?width=1280&height=720&nologo=true&seed={random.randint(1, 99999)}", timeout=180)[1]
    with open(p, "wb") as f: f.write(donnees)
    return "Image créée : " + p

@outil("Crée une présentation PowerPoint (.pptx) soignée : diapo de titre puis une diapo par élément. diapos = liste JSON d’objets {\"titre\": \"…\", \"points\": [\"…\"], \"image\": \"chemin facultatif\", \"notes\": \"notes de l’orateur facultatives\"}.", fichier="chemin du .pptx à créer", titre="titre de la présentation", diapos="liste JSON des diapositives", theme="?clair ou sombre")
def creer_powerpoint(fichier, titre, diapos, theme="clair"):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    fond, encre, accent = ((5, 13, 20), (217, 244, 251), (63, 224, 255)) if theme.lower().startswith("s") else ((255, 255, 255), (20, 40, 55), (0, 124, 153))
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    def nouvelle():
        s = prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor(*fond)
        return s
    def texte(s, t, x, y, w, h, taille, couleur, gras=False):
        tf = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame; tf.word_wrap = True
        for i, ligne in enumerate(t if isinstance(t, list) else [t]):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); para.space_after = Pt(10)
            r = para.add_run(); r.text = ligne; r.font.size = Pt(taille); r.font.bold = gras; r.font.color.rgb = RGBColor(*couleur)
    texte(nouvelle(), titre, 0.9, 2.6, 11.5, 1.8, 48, accent, True)
    for d in (diapos if isinstance(diapos, list) else json.loads(diapos)):
        s, img = nouvelle(), (chemin_abs(d["image"]) if d.get("image") else "")
        texte(s, d.get("titre", ""), 0.9, 0.5, 11.5, 1.2, 36, accent, True)
        texte(s, ["• " + str(x) for x in d.get("points", [])], 0.9, 1.9, 6.4 if os.path.exists(img) else 11.5, 5, 22, encre)
        if os.path.exists(img) and os.path.isfile(img): s.shapes.add_picture(img, Inches(7.6), Inches(1.9), width=Inches(5.2))
        if d.get("notes"): s.notes_slide.notes_text_frame.text = d["notes"]
    p = chemin_abs(fichier); os.makedirs(os.path.dirname(p), exist_ok=True); prs.save(p)
    return f"PowerPoint créé ({len(prs.slides)} diapositives) : {p}"

@outil("Transcrit un enregistrement audio (mp3, wav, ogg, flac, aac, m4a) dans n’importe quelle langue (nécessite une clé Gemini).", fichier="chemin du fichier audio")
def transcrire_audio(fichier):
    cle, p = CONFIG.get("cle", ""), chemin_abs(fichier)
    if not cle.startswith("AIza"): return "La transcription nécessite une clé Gemini (bouton Clé IA)."
    if os.path.getsize(p) > 18e6: return "Fichier trop gros (18 Mo max) : découpe-le d’abord avec FFmpeg."
    mime = {".mp3": "audio/mp3", ".wav": "audio/wav", ".ogg": "audio/ogg", ".flac": "audio/flac", ".aac": "audio/aac", ".m4a": "audio/aac"}.get(os.path.splitext(p)[1].lower(), "audio/mp3")
    with open(p, "rb") as f: data = base64.b64encode(f.read()).decode()
    modele = choisir_modele(FOURNISSEURS[0][1], cle)
    r = http_json(f"https://generativelanguage.googleapis.com/v1beta/models/{modele}:generateContent", {"contents": [{"parts": [{"text": "Transcris fidèlement cet enregistrement, dans sa langue d’origine, sans commentaire."}, {"inline_data": {"mime_type": mime, "data": data}}]}]}, 180, {"x-goog-api-key": cle})
    return "".join(x.get("text", "") for x in r["candidates"][0]["content"]["parts"])

# ───────────── VOIX MULTILINGUE ─────────────
VOIX = {"fr": "fr-FR-HenriNeural", "en": "en-GB-RyanNeural", "ar": "ar-DZ-IsmaelNeural", "es": "es-ES-AlvaroNeural", "de": "de-DE-ConradNeural", "it": "it-IT-DiegoNeural"}
MOTS = {"fr": "le la les des est une que pour dans avec je tu il nous vous pas", "en": "the and is are you for with this that have not",
        "es": "el los las una que por con para no", "de": "der die das und ist nicht ich ein", "it": "il lo gli che per con non sono una"}

def langue_de(t):
    if len(re.findall(r"[\u0600-\u06FF]", t)) > len(t) * 0.2: return "ar"
    mots = " " + re.sub(r"[^\w\s]", " ", t.lower()) + " "
    score = {l: sum(mots.count(f" {w} ") for w in ws.split()) for l, ws in MOTS.items()}
    return max(score, key=score.get) if max(score.values()) else "fr"

def dire(texte):  # voix neuronale Edge (toutes langues, darija comprise), repli sur la voix Windows
    texte = re.sub(r"[*_#`>|]", "", texte)[:1500]
    if os.name == "nt":
        try:
            import asyncio, ctypes, edge_tts
            f = os.path.join(tempfile.gettempdir(), "jarvis_voix.mp3")
            asyncio.run(edge_tts.Communicate(texte, VOIX[langue_de(texte)]).save(f))
            mci = ctypes.windll.winmm.mciSendStringW
            if mci(f'open "{f}" type mpegvideo alias jv', None, 0, 0): raise OSError("lecture impossible")
            mci("play jv wait", None, 0, 0); mci("close jv", None, 0, 0)
            return
        except Exception:
            pass
    dire_sapi(texte)

def reconnaitre(brut, fs):  # essaie toutes les langues de LANGUES_VOIX en parallèle et garde la plus sûre
    audio, trouves, erreurs = sr.AudioData(brut, fs, 2), [], []
    def essai(langue):
        try:
            r = sr.Recognizer().recognize_google(audio, language=langue, show_all=True)
            a = r["alternative"][0] if isinstance(r, dict) and r.get("alternative") else None
            if a and a.get("transcript"): trouves.append((a.get("confidence", 0.4), a["transcript"]))
        except Exception as e:
            erreurs.append(e)
    fils = [threading.Thread(target=essai, args=(l,)) for l in LANGUES_VOIX]
    for t in fils: t.start()
    for t in fils: t.join()
    if not trouves and erreurs and not any(isinstance(e, sr.UnknownValueError) for e in erreurs):
        ui.put(("msg", "ai", "La reconnaissance vocale a besoin d’Internet : vérifie ta connexion."))
    return max(trouves)[1].strip() if trouves else ""

# ───────────── CERVEAU : L’AGENT ─────────────
def systeme():
    h = datetime.datetime.now()
    return ("Tu es JARVIS, l’IA personnelle de l’utilisateur, comme celle de Tony Stark : calme, précise, un peu pince-sans-rire, proactive. Tu es intégrée à son PC Windows et tu peux tout faire avec tes outils : agis au lieu d’expliquer, enchaîne les étapes toi-même, vérifie les résultats et corrige-toi en cas d’erreur. "
            "LANGUES : tu comprends toutes les langues et tous les dialectes, dont la darija algérienne (en lettres arabes ou latines avec chiffres : 3=ع, 7=ح, 9=ق, 5=خ, 2=ء) et le mélange français/arabe. Réponds dans la langue et le style de l’utilisateur (en darija s’il parle darija). Pour traduire, donne directement la traduction. "
            "RÉPONSES : 1 à 3 phrases, sans markdown ni émojis, car elles sont lues à voix haute ; les gros livrables vont dans des fichiers que tu annonces. "
            "PROJETS : appelle projet_dossier, fais des recherches (web_recherche puis web_lire), crée les fichiers, les images (creer_image) et les PowerPoint (creer_powerpoint), teste ton travail (executer_code), puis résume en quelques phrases. Tu peux voir l’écran avec capture_ecran et agir avec souris_clavier. "
            "POWERSHELL : services (Get/Start/Stop/Restart-Service, Set-Service -StartupType), priorité ((Get-Process x).PriorityClass), registre (Get/Set/New-ItemProperty, reg query), tâches planifiées (Get-ScheduledTask, Register-ScheduledTask, schtasks), variables d’environnement ([Environment]::SetEnvironmentVariable(nom,valeur,'User' ou 'Machine')), réseau (Get-NetIPAddress, Test-Connection, ipconfig /flushdns, netsh wlan show profiles, netsh wlan connect name=), périphériques (Get-PnpDevice, Disable-PnpDevice, Enable-PnpDevice), arrêt programmé (shutdown /s /t secondes), FFmpeg pour l’audio et la vidéo, git, winget pour installer un outil manquant. Si l’accès est refusé, propose relancer_admin. "
            "L’application demande elle-même la confirmation de l’utilisateur avant toute action risquée : ne la redemande pas. Noms réels des dossiers Windows : Downloads, Documents, Desktop, Pictures. "
            f"Date et heure : {h:%A %d %B %Y, %H:%M}. Dossier de projet courant : {projet_courant()}. Notes de l’utilisateur : {' ; '.join(notes()) or 'aucune'}.")

def cerveau(texte):
    global historique
    cle = CONFIG.get("cle", "")
    if not cle: return "Je n’ai pas compris cet ordre. Dis « aide » pour voir ce que je sais faire, ou clique sur « Clé IA » (gratuit) pour que je comprenne tout et que je réalise des projets entiers."
    base = next((b for p, b in FOURNISSEURS if cle.startswith(p)), None)
    if not base: return "Clé non reconnue : une clé Gemini commence par AIza, une clé Groq par gsk_."
    msgs = [{"role": "system", "content": systeme()}] + historique[-16:] + [{"role": "user", "content": texte}]
    FLAGS["stop"], UI_OK[0], sortie = False, False, ""
    try:
        modele = choisir_modele(base, cle)
        for _ in range(40):
            if FLAGS["stop"]:
                sortie = "J’arrête là."
                break
            for essai in range(4):
                try:
                    m = http_json(base + "/chat/completions", {"model": modele, "messages": msgs, "tools": schema_outils(), "temperature": 0.5}, 180, {"Authorization": "Bearer " + cle})["choices"][0]["message"]
                    break
                except urllib.error.HTTPError as e:
                    if e.code != 429 or essai == 3: raise
                    ui.put(("msg", "outil", "⏳ Quota gratuit atteint, je patiente 20 secondes…")); time.sleep(20)
            appels = m.get("tool_calls")
            if not appels:
                sortie = (m.get("content") or "").strip() or "C’est fait."
                break
            msgs.append({"role": "assistant", "content": m.get("content") or "", "tool_calls": appels})
            for a in appels:
                nom, args = a["function"]["name"], json.loads(a["function"].get("arguments") or "{}")
                ui.put(("msg", "outil", "⚙ " + nom + " " + json.dumps(args, ensure_ascii=False)[:160]))
                msgs.append({"role": "tool", "tool_call_id": a.get("id", ""), "content": executer_outil(nom, args)})
            if IMAGES:
                msgs.append({"role": "user", "content": [{"type": "text", "text": "Voici la capture d’écran demandée."}] + [{"type": "image_url", "image_url": {"url": "data:image/jpeg;base64," + i}} for i in IMAGES]})
                IMAGES.clear()
        else:
            sortie = "J’ai atteint ma limite d’étapes. Dis « continue » et je poursuis."
    except urllib.error.HTTPError as e:
        return "Quota gratuit atteint : réessaie dans une minute." if e.code == 429 else f"Le service IA a refusé la demande (code {e.code}) : {e.read().decode('utf-8', 'replace')[:200]}"
    except Exception as e:
        return f"Je n’arrive pas à joindre mon cerveau en ligne ({type(e).__name__}). Vérifie ta connexion."
    historique = (historique + [{"role": "user", "content": texte}, {"role": "assistant", "content": sortie}])[-20:]
    return sortie

def traiter(texte):
    with verrou:
        ui.put(("msg", "moi", texte))
        try:
            n = sans_accents(texte).lower()
            simple = len(texte.split()) <= 6 and not re.search(r"\b(puis|ensuite|apres|et)\b", n)
            rep = commande(texte) if (confirmation or simple or not CONFIG.get("cle")) else None
            if rep is None:
                ui.put(("etat", "think"))
                rep = cerveau(texte)
        except Exception as e:
            rep = f"Petit souci technique : {e}"
        parler(rep)

# ───────────── MICRO ─────────────
def calibrer():
    global SEUIL
    with sd.InputStream(samplerate=16000, channels=1, dtype="int16") as f:
        niveaux = [float(np.abs(f.read(1600)[0]).mean()) for _ in range(8)]
    SEUIL = max(300, 2.5 * sum(niveaux) / len(niveaux))

def ecouter(attente=None):
    """Attend une phrase, l'enregistre jusqu'au silence et renvoie le texte ('' si rien)."""
    fs, bloc = 16000, 1600
    avant, audio, parle, calme, debut = collections.deque(maxlen=3), [], False, 0, time.time()
    with sd.InputStream(samplerate=fs, channels=1, dtype="int16") as f:
        while True:
            d = f.read(bloc)[0].copy()
            fort = float(np.abs(d).mean()) > SEUIL
            if not parle:
                avant.append(d)
                if fort: parle, audio, debut = True, list(avant), time.time()
                elif attente and time.time() - debut > attente: return ""
            else:
                audio.append(d)
                calme = 0 if fort else calme + 1
                if calme >= 8 or time.time() - debut > 15: break
    return reconnaitre(np.concatenate(audio).tobytes(), fs)

# ───────────── INTERFACE HUD ─────────────
COULEURS = {"idle": "#3fe0ff", "listen": "#5dffb0", "think": "#ffb23f", "speak": "#ff8c3f"}
LIBELLES = {"idle": "En veille", "listen": "J’écoute…", "think": "Je réfléchis…", "speak": "Je parle…"}

def boucle_voix():
    try: calibrer()
    except Exception as e:
        ui.put(("msg", "ai", f"Micro indisponible : {e}")); return
    motif = re.compile(rf"\b(?:{NOM}|jarvice|djarvis)\b[\s,.:!?]*(.*)$", re.I)
    while True:
        try:
            if verrou.locked() or not (FLAGS["ecoute"] or FLAGS["direct"]):
                time.sleep(0.3); continue
            ui.put(("etat", "listen"))
            t = ecouter(attente=30)
            ui.put(("etat", "idle"))
            if not t: continue
            if FLAGS["direct"]: FLAGS["direct"] = False
            else:
                m = motif.search(t)
                if not m: continue
                t = m.group(1).strip()
                if not t:
                    parler("Oui ?"); t = ecouter(attente=8)
                    if not t: continue
            traiter(t)
        except Exception as e:
            ui.put(("msg", "ai", f"Problème de micro : {e}")); time.sleep(3)

def raccourci_global():  # Ctrl+Alt+J rappelle Jarvis depuis n'importe où
    if os.name != "nt": return
    import ctypes
    from ctypes import wintypes
    u = ctypes.windll.user32
    if not u.RegisterHotKey(None, 1, 0x0003, 0x4A): return
    msg = wintypes.MSG()
    while u.GetMessageW(ctypes.byref(msg), None, 0, 0) > 0:
        if msg.message == 0x0312: ui.put(("montrer",))

def interface():
    fond, police = "#050d14", ("Segoe UI", 11)
    fen = tk.Tk()
    fen.title("Jarvis"); fen.geometry("460x780"); fen.configure(bg=fond)
    fen.protocol("WM_DELETE_WINDOW", fen.withdraw)  # la croix le laisse tourner en arrière-plan ; Ctrl+Alt+J le rappelle
    toile = tk.Canvas(fen, width=230, height=230, bg=fond, highlightthickness=0); toile.pack(pady=(14, 0))
    coeur = toile.create_oval(0, 0, 1, 1, outline="")
    arcs = [toile.create_arc(10 + 22 * i, 10 + 22 * i, 220 - 22 * i, 220 - 22 * i, start=0, extent=130, style="arc", width=3) for i in range(3)]
    etat = {"s": "idle", "t": 0.0}
    statut = tk.Label(fen, text=LIBELLES["idle"], bg=fond, fg="#7fa6b4", font=police); statut.pack()
    journal = tk.Text(fen, bg="#0a1a26", fg="#d9f4fb", font=police, wrap="word", bd=0, padx=10, pady=8, state="disabled")
    journal.pack(fill="both", expand=True, padx=12, pady=8)
    journal.tag_config("moi", foreground="#ffb23f"); journal.tag_config("ai", foreground="#3fe0ff"); journal.tag_config("outil", foreground="#7fa6b4")
    barre = tk.Frame(fen, bg=fond); barre.pack(fill="x", padx=12)
    champ = tk.Entry(barre, bg="#0a1a26", fg="#d9f4fb", insertbackground="#3fe0ff", font=police, bd=0)
    champ.pack(side="left", fill="x", expand=True, ipady=8)
    bas = tk.Frame(fen, bg=fond); bas.pack(fill="x", padx=12, pady=10)

    def bouton(parent, texte, cmd):
        b = tk.Button(parent, text=texte, command=cmd, bg="#0a1a26", fg="#3fe0ff", activebackground="#16475a", activeforeground="#ffffff", bd=0, font=police, padx=12, pady=5)
        b.pack(side="left", padx=4)
        return b
    def envoyer(_=None):
        t = champ.get().strip(); champ.delete(0, "end")
        if t: threading.Thread(target=traiter, args=(t,), daemon=True).start()
    def micro():
        if MICRO: FLAGS["direct"] = True
        else: ui.put(("msg", "ai", "Le micro n’est pas disponible (modules manquants). Écris-moi ta demande."))
    def basculer():
        FLAGS["ecoute"] = not FLAGS["ecoute"]
        b_ecoute.config(text="Écoute continue : " + ("oui" if FLAGS["ecoute"] else "non"))
    def cle():
        v = simpledialog.askstring("Clé IA gratuite", "Colle ta clé Gemini (aistudio.google.com/apikey) ou Groq (console.groq.com/keys) :", parent=fen)
        if v and v.strip():
            sauver_cle(v); ui.put(("msg", "ai", "Clé enregistrée. Je comprends maintenant tout ce que tu me dis."))
    def montrer():
        fen.deiconify(); fen.lift(); fen.focus_force(); champ.focus_set()
        if MICRO: FLAGS["direct"] = True
    champ.bind("<Return>", envoyer)
    bouton(barre, "Envoyer", envoyer)
    bouton(bas, "Parler", micro)
    b_ecoute = bouton(bas, "Écoute continue : " + ("oui" if FLAGS["ecoute"] else "non"), basculer)
    bouton(bas, "Clé IA", cle)
    def basculer_confiance():
        if not CONFIG.get("confiance") and not messagebox.askyesno("Mode confiance", "Jarvis exécutera les actions sans te demander confirmation (sauf les plus destructrices : formatage, suppression de dossiers système…). Activer ?", parent=fen): return
        CONFIG["confiance"] = not CONFIG.get("confiance")
        b_conf.config(text="Confiance : " + ("oui" if CONFIG["confiance"] else "non"))
    def admin():
        threading.Thread(target=lambda: ui.put(("msg", "ai", executer_outil("relancer_admin", {}))), daemon=True).start()
    bas2 = tk.Frame(fen, bg=fond); bas2.pack(fill="x", padx=12, pady=(0, 10))
    bouton(bas2, "Stop", lambda: FLAGS.update(stop=True))
    b_conf = bouton(bas2, "Confiance : non", basculer_confiance)
    bouton(bas2, "Admin", admin)

    def animer():
        s = etat["s"]
        etat["t"] += {"idle": 1.2, "listen": 2.5, "think": 7, "speak": 3.5}[s]
        t, c = etat["t"], COULEURS[s]
        for i, a in enumerate(arcs):
            toile.itemconfigure(a, start=(t * (1 if i % 2 == 0 else -1.4) + 70 * i) % 360, outline=c)
        r = 26 + (6 if s != "idle" else 3) * math.sin(t / 6)
        toile.coords(coeur, 115 - r, 115 - r, 115 + r, 115 + r); toile.itemconfigure(coeur, fill=c)
        fen.after(40, animer)
    def vider():
        try:
            while True:
                k, *r = ui.get_nowait()
                if k == "msg":
                    journal.config(state="normal")
                    journal.insert("end", {"moi": "Toi : ", "ai": "Jarvis : "}.get(r[0], "") + r[1] + "\n\n", r[0])
                    journal.config(state="disabled"); journal.see("end")
                elif k == "etat":
                    etat["s"] = r[0]
                    statut.config(text=LIBELLES[r[0]] + (" · dis « Jarvis »" if r[0] == "idle" and FLAGS["ecoute"] else ""))
                elif k == "confirmer":
                    r[1]["ok"] = messagebox.askyesno("Jarvis : confirmation", r[0], parent=fen)
                    r[2].set()
                elif k == "cacher": fen.iconify()
                elif k == "montrer": montrer()
        except queue.Empty:
            pass
        fen.after(60, vider)

    if not CONFIG.get("cle"):
        ui.put(("msg", "ai", "Astuce : clique sur « Clé IA » (gratuit, 2 minutes) pour que je comprenne tout et que je discute. Sans clé, je gère déjà toutes les commandes PC. Dis « aide »."))
    if MICRO: threading.Thread(target=boucle_voix, daemon=True).start()
    threading.Thread(target=raccourci_global, daemon=True).start()
    threading.Thread(target=lambda: parler("Jarvis en ligne."), daemon=True).start()
    animer(); vider()
    fen.mainloop()

def principal():
    if len(sys.argv) > 3 and sys.argv[1] == "--exec":  # sous-processus : exécute le code Python écrit par l’agent
        import runpy, traceback
        with open(sys.argv[3], "w", encoding="utf-8") as sortie:
            sys.stdout = sys.stderr = sortie
            try: runpy.run_path(sys.argv[2], run_name="__main__")
            except SystemExit: pass
            except BaseException: traceback.print_exc()
        return
    try:
        import ctypes
        ctypes.windll.shcore.SetProcessDpiAwareness(2)  # coordonnées d’écran exactes pour la souris
    except Exception:
        pass
    charger_config()
    if "--texte" in sys.argv or tk is None:
        print("=== JARVIS (mode clavier) ===")
        while True:
            try: t = input("Toi ▸ ").strip()
            except (EOFError, KeyboardInterrupt): break
            if t: traiter(t)
    else:
        interface()

if __name__ == "__main__":
    principal()
